import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

import Roy


class RoyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.docx"
        self.output = Path(self.temp.name) / "output.docx"

    def make_document(self, count=120, punctuation=True, size=11):
        doc = Document()
        doc.styles["Normal"].font.size = Pt(size)
        for _ in range(count):
            p = doc.add_paragraph()
            p.add_run("Przykładowy").bold = True
            p.add_run(" tekst dokumentu" + ("." if punctuation else ""))
        doc.save(self.source)
        return doc

    def test_roundtrip_polish_and_formatting(self):
        self.make_document(count=200, size=13)
        message = " Zażółć gęślą jaźń!\n "
        Roy.encode(self.source, message, self.output)
        self.assertEqual(Roy.decode(self.output), message)
        original = Document(self.source)
        encoded = Document(self.output)
        pairs = set()
        for before, after in zip(original.paragraphs, encoded.paragraphs):
            self.assertEqual(before.runs[0].bold, after.runs[0].bold)
            if len(pairs) < 4:
                try:
                    pairs.add(Roy._decode_pair(encoded, after))
                except ValueError:
                    pass
        self.assertEqual(pairs, {0, 1, 2, 3})
        self.assertEqual(original.paragraphs[-1].text, encoded.paragraphs[-1].text)

    def test_no_punctuation(self):
        self.make_document(punctuation=False, size=9)
        Roy.encode(self.source, "Wiadomość", self.output)
        self.assertEqual(Roy.decode(self.output), "Wiadomość")
        self.assertTrue(any(p.text.endswith(" .") for p in Document(self.output).paragraphs))

    def test_capacity_and_untouched_source(self):
        self.make_document(count=51)
        with self.assertRaisesRegex(ValueError, "potrzeba 52"):
            Roy.encode(self.source, "a", self.output)
        self.assertFalse(self.output.exists())
        with self.assertRaisesRegex(ValueError, "inny plik"):
            Roy.encode(self.source, "a", self.source)

    def test_invalid_carriers(self):
        doc = self.make_document()
        cases = ["dwa\tsłowa"]
        for text in cases:
            with self.subTest(text=text):
                doc.paragraphs[0].text = text
                doc.save(self.source)
                with self.assertRaisesRegex(ValueError, "Akapit 1"):
                    Roy.encode(self.source, "a", self.output)
                self.assertFalse(self.output.exists())
        doc.paragraphs[0].text = "dwa słowa"
        doc.paragraphs[0].runs[0].font.size = Pt(11)
        doc.save(self.source)
        with self.assertRaisesRegex(ValueError, "stylu akapitu"):
            Roy.encode(self.source, "a", self.output)

    def test_skips_single_word_paragraphs_for_encoding_and_decoding(self):
        doc = Document()
        single_word_texts = ["Tytuł", "Szczepan54", "Koniec ."]
        for i in range(60):
            if i in (0, 20, 48):
                doc.add_paragraph(single_word_texts.pop(0))
            doc.add_paragraph("Zwykły tekst dokumentu.")
        doc.save(self.source)

        Roy.encode(self.source, "a", self.output)

        self.assertEqual(Roy.decode(self.output), "a")
        before = Document(self.source)
        after = Document(self.output)
        for index in (0, 21, 50):
            self.assertEqual(after.paragraphs[index].text, before.paragraphs[index].text)

    def test_single_word_paragraphs_do_not_count_as_capacity(self):
        doc = self.make_document(count=51)
        for _ in range(20):
            doc.add_paragraph("Nagłówek")
        doc.save(self.source)
        with self.assertRaisesRegex(ValueError, "potrzeba 52, dostępnych 51"):
            Roy.encode(self.source, "a", self.output)

    def test_normalizes_spaces_across_formatted_runs(self):
        doc = Document()
        for _ in range(60):
            p = doc.add_paragraph()
            p.add_run("  Ala ").bold = True
            p.add_run("  ma  kota  .  ").italic = True
        doc.save(self.source)
        source_before = self.source.read_bytes()

        Roy.encode(self.source, "a", self.output)

        self.assertEqual(Roy.decode(self.output), "a")
        self.assertEqual(self.source.read_bytes(), source_before)
        result = Document(self.output)
        # The first pair in the header is 01: one space is reinserted.
        self.assertEqual(result.paragraphs[0].text, "Ala  ma kota.")
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].italic)
        # Paragraphs outside the payload must not be normalized.
        self.assertEqual(result.paragraphs[-1].text, "  Ala   ma  kota  .  ")

    def test_nonstego_and_corruption(self):
        self.make_document()
        with self.assertRaises(ValueError):
            Roy.decode(self.source)
        Roy.encode(self.source, "Test", self.output)
        doc = Document(self.output)
        # Change a valid data symbol to a different valid symbol.
        p = doc.paragraphs[48]
        old_pair = Roy._decode_pair(doc, p)
        for run in p.runs:
            run.font.size = None
        p.text = "Podwójny  odstęp" if old_pair != 1 else "Odstęp przed kropką ."
        doc.save(self.output)
        with self.assertRaisesRegex(ValueError, "kontrolna"):
            Roy.decode(self.output)

    def test_character_style_size_is_rejected(self):
        doc = self.make_document()
        style = doc.styles.add_style("Own Size", WD_STYLE_TYPE.CHARACTER)
        style.font.size = Pt(14)
        doc.paragraphs[0].runs[0].style = style
        doc.save(self.source)
        with self.assertRaisesRegex(ValueError, "Styl znaku"):
            Roy.encode(self.source, "a", self.output)


if __name__ == "__main__":
    unittest.main()
